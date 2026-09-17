import hashlib
import json
import pytest
from pathlib import Path

from site_safety.agents.knowledge_base import KnowledgeBase, _read_pages, _document_clauses
from site_safety.agents.schemas import DetectionEvent, KnowledgeReference
from test_agents import _event, _finding, _reasoning
from site_safety.agents import CollaborativeResponseAgent, ReviewLearningAgent


def test_evidence_survives_business_storage_and_confirmation(tmp_path):
    import app_server as server
    finding = _finding('missing_helmet', verified=True, confidence=.9, review=True)
    finding.knowledge_references = [KnowledgeReference(chunk_id='KB-test', source_file='real.md', text='完整条文', document_sha256='abc')]
    finding.knowledge_status = 'retrieved_unconfirmed'
    event = _event([finding])
    db = server.Database(tmp_path / 'test.db')
    try:
        state = server.AppState(db, tmp_path)
        result = state.ingest_event(event.model_dump())
        assert result['work_orders'] == 0
        assert result['confirmations'] == 1
        stored = DetectionEvent.model_validate(db.get(event.event_id))
        assert stored.risks[0].knowledge_references == finding.knowledge_references
        assert stored.risks[0].risk_level == 'pending_review'
        request = db.list('confirmation')[0]
        assert request['knowledge_references'][0]['text'] == '完整条文'
        state.decide_confirmation(request['request_id'], {'verdict': 'confirmed'}, 'safety')
        order = db.list('work_order')[0]
        assert order['knowledge_references'][0]['document_sha256'] == 'abc'
        assert db.get(event.event_id)['risks'][0]['human_review_status'] == 'confirmed'
    finally:
        db.conn.close()


def test_catalog_hash_binding_and_full_text(tmp_path):
    content = ('施工人员安全防护要求。' * 50).encode()
    (tmp_path / '规范.md').write_bytes(content)
    catalog = {'规范.md': {'document_sha256': hashlib.sha256(content).hexdigest(), 'source_status': 'official_text_checked', 'version': 'test'}}
    (tmp_path / 'source_catalog.json').write_text(json.dumps(catalog), encoding='utf-8')
    kb = KnowledgeBase(tmp_path)
    hit = kb.search('施工人员安全防护')[0]
    assert len(hit['text']) > 400
    assert hit['source_status'] == 'official_text_checked'
    kb.import_document('规范.md', '施工人员安全防护新规定，未经核验。'.encode())
    assert kb.search('施工人员')[0]['source_status'] == 'unverified'
    assert (tmp_path / '.history' / hashlib.sha256(content).hexdigest() / '规范.md').read_bytes() == content


def test_docx_table_is_read(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('施工管理要求')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格中的临边防护条文'
    path = tmp_path / '规范.docx'
    doc.save(path)
    assert '表格中的临边防护条文' in _read_pages(path)[0][1]


def test_curated_sources_and_no_match():
    kb = KnowledgeBase(Path(__file__).resolve().parents[1] / 'knowledge_base')
    hits = kb.search('劳动防护用品安全防护用具')
    assert hits
    assert all(h['source_status'] == 'official_text_checked' for h in hits)
    assert not any(c.source_file in {'页面业务逻辑与新增规范.md', '示例-高处作业安全带使用规程.md'} for c in kb.chunks)
    assert kb.search('zzqqxx987654321', min_score=.5) == []


def test_review_revision_history_and_revocation(tmp_path):
    kb = KnowledgeBase(tmp_path)
    kb.import_document('真实条文.txt', '第一条 临边防护设施必须定期检查，防止坠落事故。'.encode())
    preview = kb.preview('真实条文.txt')
    payload = dict(document_sha256=preview['document_sha256'], revision='', title='测试来源',
                   version='2026', source_url='https://example.gov.cn/standard', reason='测试人工核对',
                   effective_status='active_as_checked', valid_until='2099-01-01')
    reviewed = kb.review_source('真实条文.txt', payload, 'admin')
    assert reviewed['metadata']['reviewed_by'] == 'admin'
    assert reviewed['history'][0]['reason'] == '测试人工核对'
    with pytest.raises(ValueError, match='已变化'):
        kb.review_source('真实条文.txt', payload, 'admin')
    payload.update(revision=reviewed['revision'], effective_status='repealed')
    kb.review_source('真实条文.txt', payload, 'admin')
    assert kb.search('临边防护') == []
    assert len(kb.preview('真实条文.txt')['history']) == 2


def test_multiple_risks_do_not_share_human_verdict(tmp_path):
    event = _event([_finding('a', verified=True, confidence=.9), _finding('b', verified=True, confidence=.9)])
    agent = ReviewLearningAgent(tmp_path / 'cases.jsonl')
    agent.ingest_confirmation(event, 'a', 'confirmed', reviewer='safety')
    assert event.review.status == 'pending'
    assert event.risks[1].human_review_status == 'pending'
    agent.ingest_confirmation(event, 'b', 'rejected', reviewer='safety')
    assert event.review.status == 'confirmed'  # all reviewed, at least one actual risk


def test_review_api_permission_report_and_mobile_evidence(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    import app_server as server
    db = server.Database(tmp_path / 'business.db')
    db.seed_users()
    state = server.AppState(db, tmp_path)
    state.kb = KnowledgeBase(tmp_path / 'kb')
    state.kb.import_document('rule.txt', '第一条 临边防护设施应保持完整可靠。'.encode())
    monkeypatch.setattr(server, 'STATE', state)
    client = TestClient(server.app)
    def headers(name):
        token = client.post('/api/auth/login', json={'username':name,'password':name+'123'}).json()['token']
        return {'Authorization':'Bearer '+token}
    try:
        safety, admin = headers('safety'), headers('admin')
        preview = client.get('/api/knowledge/preview', params={'filename':'rule.txt'}, headers=admin).json()
        payload = dict(filename='rule.txt', document_sha256=preview['document_sha256'], revision='',
                       title='测试', version='2026', source_url='https://example.gov.cn/rule',
                       reason='测试核验流程', valid_until='2099-01-01', effective_status='unverified')
        assert client.post('/api/knowledge/review', json=payload, headers=safety).status_code == 403
        assert client.post('/api/knowledge/review', json=payload, headers=admin).status_code == 200
        assert client.get('/api/knowledge/source', params={'filename':'../app_server.py'}, headers=admin).status_code == 400
        assert client.get('/api/knowledge/source', params={'filename':'rule.txt'}, headers=admin).status_code == 200
        finding = _finding('missing_helmet', verified=True, confidence=.9)
        finding.knowledge_references = [KnowledgeReference.model_validate(state.kb.search('临边防护')[0])]
        state.ingest_event(_event([finding]).model_dump())
        report = client.get('/api/events/EVT-TEST-00000001/evidence-report', headers=safety).json()
        assert '临边防护设施应保持完整可靠' in report['markdown']
        assert 'SHA256' in report['markdown']
        alarm = client.get('/api/mobile/alarm/latest', headers=safety).json()
        assert alarm['knowledgeReferences'][0]['reviewed_by'] == 'admin'
        assert alarm['humanReviewStatus'] == 'pending'
        state.kb.delete_document('rule.txt')
        assert client.get('/api/events/EVT-TEST-00000001/evidence-report', headers=safety).json() == report
    finally:
        db.conn.close()


def test_upload_preview_does_not_import(tmp_path):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    from detect_bridge import create_app
    kb = KnowledgeBase(tmp_path)
    client = TestClient(create_app(SimpleNamespace(knowledge_base=kb)))
    response = client.post('/api/detect/knowledge/preview-upload', files={'file':('new.txt', '临边防护设施应保持完整可靠。'.encode())})
    assert response.status_code == 200
    assert response.json()['chunks']
    assert not (tmp_path / 'new.txt').exists()
    assert kb.chunks == []
    assert client.post('/api/detect/knowledge/preview-upload', files={'file':('bad.pdf', b'bad')}).status_code == 422


def test_expired_source_requires_review_and_pdf_garbage_rejected(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import pypdf
    content = '临边防护设施应保持完整可靠。'.encode()
    (tmp_path / 'rule.txt').write_bytes(content)
    (tmp_path / 'source_catalog.json').write_text(json.dumps({'rule.txt':{
        'document_sha256':hashlib.sha256(content).hexdigest(), 'source_status':'human_checked',
        'effective_status':'active_as_checked', 'valid_until':'2000-01-01'}}), encoding='utf-8')
    assert KnowledgeBase(tmp_path).search('临边防护')[0]['effective_status'] == 'review_due'
    monkeypatch.setattr(pypdf, 'PdfReader', lambda path: SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda:'\x00'*30+'乱码')]))
    with pytest.raises(ValueError, match='乱码'):
        _read_pages(tmp_path / 'fake.pdf')


def test_pdf_cross_page_clause_retains_exception_and_page_provenance():
    chunks = _document_clauses([(6, '3.2.1 高处作业应采取防护措施。\n以下情况需结合现场判断：'),
                               (7, '特殊条件不得仅凭图像判定。\n3.2.2 临边防护应保持完整可靠。')])
    assert len(chunks) == 2
    assert '特殊条件不得仅凭图像判定' in chunks[0]['text']
    assert chunks[0]['page_numbers'] == [6, 7]
    assert chunks[1]['page_numbers'] == [7]
    assert '3.2.2' not in chunks[0]['text']


def test_unstructured_pdf_does_not_invent_cross_page_clauses():
    chunks = _document_clauses([(1, '这是无编号的第一页安全说明。'), (2, '这是完全不同的第二页说明。')])
    assert len(chunks) == 2
    assert chunks[0]['page_numbers'] == [1]


@pytest.mark.parametrize('status', ['repealed', 'superseded'])
def test_expired_revocation_never_reenters_search(tmp_path, status):
    content = '安全防护条文，临边设施检查要求。'.encode()
    (tmp_path / 'rule.txt').write_bytes(content)
    (tmp_path / 'source_catalog.json').write_text(json.dumps({'rule.txt': {
        'document_sha256': hashlib.sha256(content).hexdigest(),
        'effective_status': status, 'valid_until': '2000-01-01'}}), encoding='utf-8')
    assert KnowledgeBase(tmp_path).search('安全防护') == []
